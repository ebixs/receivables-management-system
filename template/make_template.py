"""
ساخت فایل قالب خالی برای ورود داده‌ی واقعی.
Builds the blank input template.

    python make_template.py      →  input_template.xlsx

شیت‌ها:
  Guide                        راهنمای ستون‌ها (موتور به آن دست نمی‌زند)
  Main sheet                   فهرست پیگیری شما (Excel Table خالی)
  Initial raw financial sheet  محل چسباندن خروجی خام بخش مالی (Excel Table خالی)
"""
import os

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

COLUMNS = ["counter", "term", "telephone", "date", "price_of_mounth",
           "state", "name_family", "status", "coach"]

# (ستون، توضیح فارسی، English, نمونه، اجباری؟)
COLUMN_GUIDE = [
    ("counter", "شماره ردیف — خالی بگذارید؛ موتور خودش ۱، ۲، ۳، … می‌نویسد",
     "Row number — leave empty, the engine fills it", "", "خیر"),
    ("term", "شناسه‌ی یکتای قسط یا چک (هر قسط یک term جدا؛ تکراری نباشد)",
     "Unique id of the installment / cheque", "500123", "بله"),
    ("telephone", "شماره تماس، فقط رقم", "Phone number, digits only", "9121234567", "خیر"),
    ("date", "تاریخ سررسید قسط به شمسی", "Due date (Jalali)", "1405/08/15", "بله"),
    ("price_of_mounth", "مبلغ قسط به ریال، فقط عدد", "Installment amount in rials", "40000000", "بله"),
    ("state", "مقطع یا نوع خدمت (آزاد)", "Program / service type (free text)", "کارشناسی", "خیر"),
    ("name_family", "نام و نام خانوادگی — همیشه به یک شکل بنویسید", "Full name — always spelled the same way", "علی نمونه", "بله"),
    ("status", "خالی = عادی · «قرمز» = قرمز · هر چیز دیگر = نام نمایندگی",
     "Empty = normal · قرمز = red · anything else = agency name", "agency 1", "خیر"),
    ("coach", "نام کارشناس مسئول (برای نمایندگی‌ها نام مدیر؛ برای قرمز خالی)",
     "Responsible expert (Manager for agencies, empty for red)", "Expert 1", "خیر"),
]

STEPS = [
    "۱. اسم شیت‌ها و اسم ستون‌ها را عوض نکنید.",
    "۲. در «Main sheet» فهرست فعلی افرادی که پیگیری می‌کنید را وارد کنید (یک ردیف برای هر قسط). "
    "داده را از ردیف ۲ داخل همان جدول بنویسید یا بچسبانید؛ جدول خودش بزرگ می‌شود.",
    "۳. هر بار قبل از اجرا، آخرین خروجی بخش مالی را در «Initial raw financial sheet» بچسبانید "
    "(ردیف‌های قبلی آن را پاک کنید). این شیت هرگز توسط برنامه تغییر نمی‌کند.",
    "۴. در اپ دسکتاپ «پردازش فایل اکسل جدید» را بزنید و همین فایل را انتخاب کنید. "
    "خروجی در پوشه‌ی Excel کنار همین فایل ساخته می‌شود.",
    "۵. دفعه‌ی بعد، همان فایل خروجی را ورودی بگیرید: فقط شیت خام را با خروجی جدید مالی عوض کنید.",
    "نکته: اگر Main sheet را خالی شروع کنید، همه‌ی ردیف‌های شیت خام «جدید» حساب می‌شوند. "
    "کارشناس هر نفر را یک‌بار دستی در ستون coach بنویسید تا موتور از دفعه‌ی بعد خودش تخصیص دهد.",
]

HEADER_FILL = PatternFill("solid", fgColor="2E75B6")
TITLE_FILL = PatternFill("solid", fgColor="1F4E78")
THIN = Side(style="thin", color="BFBFBF")


def _data_sheet(wb, title, table_name, note):
    ws = wb.create_sheet(title)
    for c, name in enumerate(COLUMNS, 1):
        cell = ws.cell(1, c, name)
        cell.font = Font(bold=True, color="FFFFFF", name="Tahoma")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        guide = next(g for g in COLUMN_GUIDE if g[0] == name)
        cell.comment = Comment(f"{guide[1]}\n\n{guide[2]}\nنمونه: {guide[3]}", "Guide")
        ws.column_dimensions[get_column_letter(c)].width = 18
    ws.column_dimensions["G"].width = 26
    # یک ردیف خالی؛ Table اکسل بدون حداقل یک ردیف داده معتبر نیست
    for c in range(1, len(COLUMNS) + 1):
        ws.cell(2, c).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(2, 3).number_format = "0"
    ws.cell(2, 4).number_format = "@"
    ws.cell(2, 5).number_format = "#,##0"
    table = Table(displayName=table_name, ref=f"A1:{get_column_letter(len(COLUMNS))}2")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "A2"
    ws.cell(1, 1).comment = Comment(note + "\n\n" + ws.cell(1, 1).comment.text, "Guide")
    return ws


def build(path):
    wb = Workbook()
    guide = wb.active
    guide.title = "Guide"
    guide.sheet_view.rightToLeft = True

    guide["A1"] = "راهنمای فایل ورودی ماشین وصول مطالبات · Input template guide"
    guide["A1"].font = Font(bold=True, size=16, color="FFFFFF", name="Tahoma")
    guide["A1"].fill = TITLE_FILL
    guide.merge_cells("A1:E1")
    guide.row_dimensions[1].height = 30

    r = 3
    guide.cell(r, 1, "مراحل").font = Font(bold=True, size=12, name="Tahoma")
    for step in STEPS:
        r += 1
        guide.cell(r, 1, step).alignment = Alignment(wrap_text=True, vertical="center",
                                                     horizontal="right", readingOrder=2)
        guide.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        guide.row_dimensions[r].height = 40

    r += 2
    guide.cell(r, 1, "ستون‌ها").font = Font(bold=True, size=12, name="Tahoma")
    r += 1
    for c, h in enumerate(["ستون", "توضیح", "English", "نمونه", "اجباری"], 1):
        cell = guide.cell(r, c, h)
        cell.font = Font(bold=True, color="FFFFFF", name="Tahoma")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in COLUMN_GUIDE:
        r += 1
        for c, v in enumerate(row, 1):
            cell = guide.cell(r, c, v)
            cell.alignment = Alignment(wrap_text=True, vertical="center",
                                       horizontal="center" if c in (1, 4, 5) else "right")
            cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
        guide.row_dimensions[r].height = 32

    for col, width in zip("ABCDE", (18, 60, 44, 16, 10)):
        guide.column_dimensions[col].width = width

    _data_sheet(wb, "Main sheet", "tbl_main",
                "این شیت: فهرست پیگیری شما؛ از ردیف ۲ داخل جدول وارد کنید.")
    _data_sheet(wb, "Initial raw financial sheet", "tbl_raw",
                "این شیت: آخرین خروجی خام بخش مالی را اینجا بچسبانید.")
    wb.save(path)
    print("written:", path)


if __name__ == "__main__":
    build(os.path.join(os.path.dirname(os.path.abspath(__file__)), "input_template.xlsx"))
