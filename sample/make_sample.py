"""
ساخت فایل نمونه‌ی کاملاً ساختگی برای امتحان برنامه.
Builds a fully fictional sample input (no real people, phones or amounts).

    python make_sample.py            →  sample_input.xlsx

شیت‌ها:
  Main sheet                   فهرست اصلی پیگیری (Excel Table)
  Initial raw financial sheet  خروجی خام مالی
کارشناس‌ها: Expert 1..3  ·  مدیر: Manager  ·  نمایندگی‌ها: agency 1, agency 2  ·  وضعیت قرمز: قرمز
"""
import os
import random

from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

COLUMNS = ["counter", "term", "telephone", "date", "price_of_mounth",
           "state", "name_family", "status", "coach"]
EXPERTS = ["Expert 1", "Expert 2", "Expert 3"]
MANAGER = "Manager"
AGENCIES = ["agency 1", "agency 2"]
RED = "قرمز"


def build(path, seed=1405):
    rnd = random.Random(seed)
    first = ["علی", "مریم", "رضا", "زهرا", "حسین", "فاطمه", "محمد", "سارا",
             "امیر", "نرگس", "مهدی", "لیلا", "حامد", "نازنین", "پویا"]
    last = ["نمونه‌زاده", "آزمایشی", "فرضی‌پور", "ساختگی", "تستی", "مثالی",
            "نمونه‌ای", "فرضیان", "آزمونی", "الگویی"]
    people = []
    while len(people) < 45:
        name = f"{rnd.choice(first)} {rnd.choice(last)}"
        if name not in people:
            people.append(name)

    kind = {}
    for i, p in enumerate(people):
        kind[p] = "red" if i < 5 else (AGENCIES[i % 2] if i < 12 else "normal")
    coach_of = {p: rnd.choice(EXPERTS) for p in people}
    phone = {p: 9000000000 + rnd.randint(0, 99_999_999) for p in people}

    def row(term, p, status, coach):
        return [None, term, phone.get(p, 9000000000 + rnd.randint(0, 99_999_999)),
                f"1405/{rnd.randint(1, 12):02d}/{rnd.randint(1, 29):02d}",
                rnd.choice([20, 30, 40, 50]) * 1_000_000,
                rnd.choice(["کاردانی", "کارشناسی", "دیپلم"]), p, status, coach]

    main, term = [], 500000
    for _ in range(140):
        p = rnd.choice(people)
        term += rnd.randint(1, 7)
        k = kind[p]
        status = RED if k == "red" else ("" if k == "normal" else k)
        coach = "" if k == "red" else (MANAGER if k != "normal" else coach_of[p])
        main.append(row(term, p, status, coach))

    # خروجی خام مالی: همه‌ی termها به‌جز چک‌های پاس‌شده + اقساط جدید
    cleared = {r[1] for r in rnd.sample(main, 10)}
    raw = []
    for r in main:
        if r[1] in cleared:
            continue
        rr = list(r)
        rr[7] = rr[8] = ""                      # خروجی مالی معمولاً status/coach ندارد
        if rnd.random() < 0.05:
            rr[2] += 1                          # مغایرت تلفن
        if rnd.random() < 0.03:
            rr[4] += 1_000_000                  # مغایرت مبلغ
        raw.append(rr)

    newcomers = [f"تازه‌وارد نمونه {i}" for i in range(1, 6)]
    forced = people[:2] + people[5:7] + people[12:14]   # قرمز، نمایندگی، عادی با سابقه
    for p in forced + newcomers + newcomers[:2] + [rnd.choice(people) for _ in range(15)]:
        term += rnd.randint(1, 7)
        raw.append(row(term, p, "", ""))

    wb = Workbook()
    ws = wb.active
    ws.title = "Main sheet"
    ws.append(COLUMNS)
    for r in main:
        ws.append(r)
    table = Table(displayName="tbl_main", ref=f"A1:I{len(main) + 1}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1")
    ws.add_table(table)

    ws_raw = wb.create_sheet("Initial raw financial sheet")
    ws_raw.append(COLUMNS)
    for r in raw:
        ws_raw.append(r)
    wb.save(path)
    print("written:", path)


if __name__ == "__main__":
    build(os.path.join(os.path.dirname(os.path.abspath(__file__)), "sample_input.xlsx"))
